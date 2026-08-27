"""Tests for document text extraction.

Prints only safe metadata (ids, filenames, page counts, char counts).
Does not send extracted text to an LLM.
"""

from __future__ import annotations

import logging
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from app.ingestion.loaders import (
    NO_READABLE_TEXT,
    ExtractedDocument,
    ExtractionError,
    load_document,
)

FIXTURES = Path(__file__).resolve().parent / "fixtures"
PAPER_PDF = FIXTURES / "attention_is_all_you_need.pdf"
PAPER_URL = "https://arxiv.org/pdf/1706.03762"


def _log_safe_metadata(document: ExtractedDocument) -> None:
    char_counts = [len(page.text) for page in document.pages]
    print(
        "extraction_ok "
        f"document_id={document.document_id} "
        f"filename={document.source_filename} "
        f"file_type={document.file_type} "
        f"page_count={document.page_count} "
        f"total_chars={sum(char_counts)} "
        f"chars_per_page={char_counts}"
    )


@pytest.fixture(scope="session")
def research_paper_pdf() -> Path:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    if PAPER_PDF.exists() and PAPER_PDF.stat().st_size > 10_000:
        return PAPER_PDF

    import urllib.request

    request = urllib.request.Request(
        PAPER_URL,
        headers={"User-Agent": "trust-aware-rag-tests/1.0"},
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            data = response.read()
    except Exception as exc:
        pytest.fail(f"Failed to download research paper PDF from {PAPER_URL}: {exc}")

    if len(data) < 10_000 or not data.startswith(b"%PDF"):
        pytest.fail("Downloaded research paper is not a valid PDF.")

    PAPER_PDF.write_bytes(data)
    return PAPER_PDF


def test_pdf_research_paper_returns_text_and_pages(research_paper_pdf: Path) -> None:
    document_id = "doc-attention-paper"
    extracted = load_document(
        research_paper_pdf,
        document_id=document_id,
        source_filename="attention_is_all_you_need.pdf",
    )
    _log_safe_metadata(extracted)

    assert extracted.document_id == document_id
    assert extracted.source_filename == "attention_is_all_you_need.pdf"
    assert extracted.file_type == "pdf"
    assert extracted.page_count >= 2

    page_numbers = [page.page_number for page in extracted.pages]
    assert page_numbers == list(range(1, extracted.page_count + 1))

    for page in extracted.pages:
        assert page.document_id == document_id
        assert page.source_filename == "attention_is_all_you_need.pdf"
        assert isinstance(page.page_number, int)
        assert page.page_number >= 1
        assert page.text.strip()
        assert len(page.text) > 20

    first_page = extracted.pages[0].text.lower()
    assert "attention" in first_page
    assert "transformer" in first_page or "sequence" in first_page


def test_pdf_extraction_logs_metadata_not_full_text(
    research_paper_pdf: Path, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level(logging.INFO, logger="app.ingestion.loaders"):
        extracted = load_document(
            research_paper_pdf,
            document_id="doc-log-check",
            source_filename="attention_is_all_you_need.pdf",
        )

    joined = " ".join(record.getMessage() for record in caplog.records)
    assert "document_id=doc-log-check" in joined
    assert "filename=attention_is_all_you_need.pdf" in joined
    assert f"page_count={extracted.page_count}" in joined
    assert extracted.pages[0].text not in joined


def test_pdf_without_text_returns_clear_error(tmp_path: Path) -> None:
    blank = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with blank.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(ExtractionError, match=NO_READABLE_TEXT):
        load_document(blank, document_id="doc-blank", source_filename="blank.pdf")


def test_txt_uses_form_feed_as_page_breaks(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("First section\fSecond section", encoding="utf-8")

    extracted = load_document(path, document_id="doc-txt", source_filename="notes.txt")
    _log_safe_metadata(extracted)

    assert [page.page_number for page in extracted.pages] == [1, 2]
    assert extracted.pages[0].text == "First section"
    assert extracted.pages[1].text == "Second section"


def test_empty_txt_raises(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("   \n", encoding="utf-8")
    with pytest.raises(ExtractionError, match=NO_READABLE_TEXT):
        load_document(path, document_id="doc-empty-txt")


def test_docx_preserves_paragraph_text(tmp_path: Path) -> None:
    path = tmp_path / "memo.docx"
    doc = Document()
    doc.add_paragraph("Trust-aware retrieval")
    doc.add_paragraph("Hallucination-resistant answers")
    doc.save(path)

    extracted = load_document(path, document_id="doc-docx", source_filename="memo.docx")
    _log_safe_metadata(extracted)

    assert extracted.pages[0].document_id == "doc-docx"
    assert extracted.pages[0].source_filename == "memo.docx"
    assert extracted.pages[0].page_number == 1
    combined = extracted.pages[0].text
    assert "Trust-aware retrieval" in combined
    assert "Hallucination-resistant answers" in combined


def test_csv_uses_row_section_identifiers(tmp_path: Path) -> None:
    path = tmp_path / "rows.csv"
    path.write_text("name,score\nalpha,1\nbeta,2\n", encoding="utf-8")

    extracted = load_document(path, document_id="doc-csv", source_filename="rows.csv")
    _log_safe_metadata(extracted)

    assert [page.page_number for page in extracted.pages] == ["row:1", "row:2"]
    assert "name: alpha" in extracted.pages[0].text
    assert "score: 2" in extracted.pages[1].text


def test_json_uses_key_section_identifiers(tmp_path: Path) -> None:
    path = tmp_path / "paper.json"
    path.write_text(
        '{"title": "Attention Is All You Need", "year": 2017}',
        encoding="utf-8",
    )

    extracted = load_document(path, document_id="doc-json", source_filename="paper.json")
    _log_safe_metadata(extracted)

    sections = {page.page_number: page.text for page in extracted.pages}
    assert sections["section:title"] == "Attention Is All You Need"
    assert "2017" in sections["section:year"]


def test_invalid_json_raises(tmp_path: Path) -> None:
    path = tmp_path / "broken.json"
    path.write_text("{not json", encoding="utf-8")
    with pytest.raises(ExtractionError, match="Failed to parse JSON"):
        load_document(path, document_id="doc-bad-json")


def test_invalid_doc_raises(tmp_path: Path) -> None:
    path = tmp_path / "fake.doc"
    path.write_bytes(b"this is not an ole document")
    with pytest.raises(ExtractionError, match="not a valid OLE Word document"):
        load_document(path, document_id="doc-bad-doc")


def test_unsupported_extension_raises(tmp_path: Path) -> None:
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG")
    with pytest.raises(ExtractionError, match="Unsupported file type"):
        load_document(path, document_id="doc-png")


def test_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(ExtractionError, match="File not found"):
        load_document(tmp_path / "missing.pdf", document_id="doc-missing")
