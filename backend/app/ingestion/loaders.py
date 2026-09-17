"""Document text extraction loaders.

Extracted text is returned as structured page/section records for later
chunking and indexing. This module does not call any LLM.
"""

from __future__ import annotations

import csv
import json
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable, Union

import olefile
from pypdf import PdfReader
from pypdf.errors import PdfReadError

logger = logging.getLogger(__name__)

PageId = Union[int, str]

NO_READABLE_TEXT = "No readable text was found in this document."
UNSUPPORTED_TYPE = "Unsupported file type"


class ExtractionError(Exception):
    """Raised when document text cannot be extracted."""


@dataclass(frozen=True)
class ExtractedPage:
    document_id: str
    source_filename: str
    page_number: PageId
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    document_id: str
    source_filename: str
    file_type: str
    pages: tuple[ExtractedPage, ...]

    @property
    def page_count(self) -> int:
        return len(self.pages)


def load_document(
    file_path: str | Path,
    document_id: str,
    source_filename: str | None = None,
) -> ExtractedDocument:
    """Dispatch to a format-specific loader.

    ``source_filename`` should be the original upload name. If omitted, the
    on-disk basename is used.
    """
    if isinstance(file_path, str) and (file_path.startswith("http://") or file_path.startswith("https://")):
        return load_url(file_path, document_id=document_id)

    path = Path(file_path)
    if not path.exists():
        raise ExtractionError(f"File not found: {path.name}")
    if not path.is_file():
        raise ExtractionError(f"Not a file: {path.name}")

    filename = source_filename or path.name
    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise ExtractionError(f"{UNSUPPORTED_TYPE}: '{suffix or path.name}'")

    try:
        pages = loader(path, document_id, filename)
    except ExtractionError:
        logger.error(
            "Extraction failed document_id=%s filename=%s file_type=%s",
            document_id,
            filename,
            suffix.lstrip("."),
        )
        raise
    except Exception as exc:
        logger.error(
            "Extraction failed document_id=%s filename=%s file_type=%s error_type=%s",
            document_id,
            filename,
            suffix.lstrip("."),
            type(exc).__name__,
        )
        raise ExtractionError(
            f"Failed to extract text from '{filename}': {type(exc).__name__}"
        ) from exc

    document = ExtractedDocument(
        document_id=document_id,
        source_filename=filename,
        file_type=suffix.lstrip("."),
        pages=tuple(pages),
    )
    _log_extraction_summary(document)
    return document


def load_pdf(path: Path, document_id: str, source_filename: str) -> list[ExtractedPage]:
    try:
        reader = PdfReader(str(path))
    except PdfReadError as exc:
        raise ExtractionError(
            f"Failed to read PDF '{source_filename}': {exc}"
        ) from exc
    except Exception as exc:
        raise ExtractionError(
            f"Failed to open PDF '{source_filename}': {type(exc).__name__}"
        ) from exc

    if getattr(reader, "is_encrypted", False):
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:
            raise ExtractionError(
                f"Failed to decrypt PDF '{source_filename}': {type(exc).__name__}"
            ) from exc
        if not unlocked:
            raise ExtractionError(
                f"PDF '{source_filename}' is encrypted and cannot be read."
            )

    if len(reader.pages) == 0:
        raise ExtractionError(NO_READABLE_TEXT)

    pages: list[ExtractedPage] = []
    page_errors: list[str] = []

    for index, page in enumerate(reader.pages, start=1):
        try:
            extracted = page.extract_text() or ""
        except Exception as exc:
            page_errors.append(f"page {index}: {type(exc).__name__}")
            continue
        text = _normalize_text(extracted)
        if text:
            pages.append(
                ExtractedPage(
                    document_id=document_id,
                    source_filename=source_filename,
                    page_number=index,
                    text=text,
                )
            )

    if page_errors:
        raise ExtractionError(
            f"Failed to extract one or more PDF pages in '{source_filename}': "
            + "; ".join(page_errors)
        )

    if not pages:
        raise ExtractionError(NO_READABLE_TEXT)

    return pages


def load_txt(path: Path, document_id: str, source_filename: str) -> list[ExtractedPage]:
    raw = _read_text_file(path, source_filename)
    sections = raw.split("\f")
    pages: list[ExtractedPage] = []
    for index, section in enumerate(sections, start=1):
        text = _normalize_text(section)
        if not text:
            continue
        pages.append(
            ExtractedPage(
                document_id=document_id,
                source_filename=source_filename,
                page_number=index,
                text=text,
            )
        )
    if not pages:
        raise ExtractionError(NO_READABLE_TEXT)
    return pages


def load_docx(path: Path, document_id: str, source_filename: str) -> list[ExtractedPage]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except Exception as exc:
        raise ExtractionError(
            f"DOCX support is unavailable on this system: {exc}"
        ) from exc

    try:
        document = Document(str(path))
    except Exception as exc:
        raise ExtractionError(
            f"Failed to read DOCX '{source_filename}': {type(exc).__name__}"
        ) from exc

    buffers: list[list[str]] = [[]]
    try:
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                _append_paragraph_with_breaks(block, buffers)
            elif isinstance(block, Table):
                table_text = _table_to_text(block)
                if table_text:
                    buffers[-1].append(table_text)
    except Exception as exc:
        raise ExtractionError(
            f"Failed to extract DOCX text from '{source_filename}': {type(exc).__name__}"
        ) from exc

    pages = _buffers_to_pages(buffers, document_id, source_filename)
    if not pages:
        raise ExtractionError(NO_READABLE_TEXT)
    return pages


def load_doc(path: Path, document_id: str, source_filename: str) -> list[ExtractedPage]:
    try:
        text = _extract_doc_binary_text(path)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            f"Failed to extract DOC text from '{source_filename}': {type(exc).__name__}"
        ) from exc

    text = _normalize_text(text)
    if not text:
        raise ExtractionError(NO_READABLE_TEXT)

    return [
        ExtractedPage(
            document_id=document_id,
            source_filename=source_filename,
            page_number="section:1",
            text=text,
        )
    ]


def load_csv(path: Path, document_id: str, source_filename: str) -> list[ExtractedPage]:
    try:
        raw = _read_text_file(path, source_filename)
    except ExtractionError:
        raise

    try:
        dialect = csv.Sniffer().sniff(raw[:4096]) if raw.strip() else csv.excel
    except csv.Error:
        dialect = csv.excel

    try:
        reader = csv.reader(raw.splitlines(), dialect)
        rows = list(reader)
    except csv.Error as exc:
        raise ExtractionError(
            f"Failed to parse CSV '{source_filename}': {exc}"
        ) from exc

    pages: list[ExtractedPage] = []
    header = rows[0] if rows else []
    data_rows = rows[1:] if len(rows) > 1 else []

    if header and data_rows:
        for index, row in enumerate(data_rows, start=1):
            pairs = []
            for col_index, value in enumerate(row):
                column = header[col_index] if col_index < len(header) else f"column_{col_index + 1}"
                pairs.append(f"{column}: {value}")
            text = _normalize_text(" | ".join(pairs))
            if not text:
                continue
            pages.append(
                ExtractedPage(
                    document_id=document_id,
                    source_filename=source_filename,
                    page_number=f"row:{index}",
                    text=text,
                )
            )
    else:
        for index, row in enumerate(rows, start=1):
            text = _normalize_text(" | ".join(row))
            if not text:
                continue
            pages.append(
                ExtractedPage(
                    document_id=document_id,
                    source_filename=source_filename,
                    page_number=f"row:{index}",
                    text=text,
                )
            )

    if not pages:
        raise ExtractionError(NO_READABLE_TEXT)
    return pages


def load_json(path: Path, document_id: str, source_filename: str) -> list[ExtractedPage]:
    raw = _read_text_file(path, source_filename)
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ExtractionError(
            f"Failed to parse JSON '{source_filename}': {exc.msg} "
            f"(line {exc.lineno}, column {exc.colno})"
        ) from exc

    pages: list[ExtractedPage] = []

    if isinstance(payload, dict):
        if not payload:
            raise ExtractionError(NO_READABLE_TEXT)
        for key, value in payload.items():
            text = _normalize_text(_json_value_to_text(value))
            if not text:
                continue
            pages.append(
                ExtractedPage(
                    document_id=document_id,
                    source_filename=source_filename,
                    page_number=f"section:{key}",
                    text=text,
                )
            )
    elif isinstance(payload, list):
        if not payload:
            raise ExtractionError(NO_READABLE_TEXT)
        for index, item in enumerate(payload, start=1):
            text = _normalize_text(_json_value_to_text(item))
            if not text:
                continue
            pages.append(
                ExtractedPage(
                    document_id=document_id,
                    source_filename=source_filename,
                    page_number=f"item:{index}",
                    text=text,
                )
            )
    else:
        text = _normalize_text(_json_value_to_text(payload))
        if text:
            pages.append(
                ExtractedPage(
                    document_id=document_id,
                    source_filename=source_filename,
                    page_number="section:document",
                    text=text,
                )
            )

    if not pages:
        raise ExtractionError(NO_READABLE_TEXT)
    return pages


def _json_value_to_text(value: object) -> str:
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, indent=2)


def _read_text_file(path: Path, source_filename: str) -> str:
    data = path.read_bytes()
    if not data:
        raise ExtractionError(NO_READABLE_TEXT)

    encodings = ("utf-8-sig", "utf-8", "utf-16", "cp1252")
    errors: list[str] = []
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: {exc.reason}")

    raise ExtractionError(
        f"Failed to decode '{source_filename}' as text. Tried: "
        + "; ".join(errors)
    )


def _normalize_text(text: str) -> str:
    cleaned = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in cleaned.split("\n")]
    collapsed = "\n".join(lines).strip()
    return collapsed


def _append_paragraph_with_breaks(paragraph: Paragraph, buffers: list[list[str]]) -> None:
    """Append paragraph text, starting a new page when Word page breaks appear."""
    xml = paragraph._element.xml
    has_break = (
        "lastRenderedPageBreak" in xml
        or 'w:type="page"' in xml
        or "w:type='page'" in xml
    )
    if has_break and any(part.strip() for part in buffers[-1]):
        buffers.append([])

    text = paragraph.text
    if text:
        buffers[-1].append(text)


def _table_to_text(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _buffers_to_pages(
    buffers: Iterable[list[str]],
    document_id: str,
    source_filename: str,
) -> list[ExtractedPage]:
    pages: list[ExtractedPage] = []
    for index, parts in enumerate(buffers, start=1):
        text = _normalize_text("\n".join(parts))
        if not text:
            continue
        pages.append(
            ExtractedPage(
                document_id=document_id,
                source_filename=source_filename,
                page_number=index,
                text=text,
            )
        )
    return pages


def _extract_doc_binary_text(path: Path) -> str:
    if not olefile.isOleFile(str(path)):
        raise ExtractionError(
            "The uploaded .doc file is not a valid OLE Word document."
        )

    ole = olefile.OleFileIO(str(path))
    try:
        if not ole.exists("WordDocument"):
            raise ExtractionError(
                "The uploaded .doc file does not contain a WordDocument stream."
            )
        word_stream = ole.openstream("WordDocument").read()
        table_name = _doc_table_stream_name(word_stream, ole)
        table_stream = ole.openstream(table_name).read()
        return _extract_text_from_piece_table(word_stream, table_stream)
    finally:
        ole.close()


def _doc_table_stream_name(word_stream: bytes, ole: olefile.OleFileIO) -> str:
    if len(word_stream) < 12:
        raise ExtractionError("The uploaded .doc file FIB is truncated.")
    flags = int.from_bytes(word_stream[10:12], "little")
    preferred = "1Table" if flags & 0x0200 else "0Table"
    if ole.exists(preferred):
        return preferred
    alternate = "0Table" if preferred == "1Table" else "1Table"
    if ole.exists(alternate):
        return alternate
    raise ExtractionError("The uploaded .doc file is missing a Word table stream.")


def _extract_text_from_piece_table(word_stream: bytes, table_stream: bytes) -> str:
    if len(word_stream) < 34:
        raise ExtractionError("The uploaded .doc file FIB is truncated.")

    csw = int.from_bytes(word_stream[32:34], "little")
    pos = 34 + 2 * csw
    if pos + 2 > len(word_stream):
        raise ExtractionError("The uploaded .doc file FIB is truncated.")
    cslw = int.from_bytes(word_stream[pos : pos + 2], "little")
    pos += 2 + 4 * cslw
    if pos + 2 > len(word_stream):
        raise ExtractionError("The uploaded .doc file FIB is truncated.")
    pos += 2  # cbRgFcLcb

    # FibRgFcLcb97: fcClx / lcbClx is the 34th field pair (0-based index 33).
    fc_off = pos + 33 * 8
    if fc_off + 8 > len(word_stream):
        raise ExtractionError("The uploaded .doc file FIB does not contain a piece table.")

    fc_clx = int.from_bytes(word_stream[fc_off : fc_off + 4], "little")
    lcb_clx = int.from_bytes(word_stream[fc_off + 4 : fc_off + 8], "little")
    if lcb_clx <= 0 or fc_clx < 0 or fc_clx + lcb_clx > len(table_stream):
        raise ExtractionError("The uploaded .doc file does not contain a valid text piece table.")

    clx = table_stream[fc_clx : fc_clx + lcb_clx]
    pcd_bytes = _find_plcfpcd(clx)
    if pcd_bytes is None:
        raise ExtractionError("The uploaded .doc file does not contain a valid text piece table.")

    length = len(pcd_bytes)
    if length < 16 or (length - 4) % 12 != 0:
        raise ExtractionError("The uploaded .doc file piece table is malformed.")

    piece_count = (length - 4) // 12
    pieces: list[str] = []
    for index in range(piece_count):
        cp_start = int.from_bytes(pcd_bytes[index * 4 : (index + 1) * 4], "little")
        cp_end = int.from_bytes(pcd_bytes[(index + 1) * 4 : (index + 2) * 4], "little")
        if cp_end < cp_start:
            raise ExtractionError("The uploaded .doc file piece table is malformed.")
        pcd_offset = (piece_count + 1) * 4 + index * 8
        pcd = pcd_bytes[pcd_offset : pcd_offset + 8]
        fc_val = int.from_bytes(pcd[2:6], "little")
        compressed = bool(fc_val & 0x40000000)
        fc = fc_val & 0x3FFFFFFF
        char_count = cp_end - cp_start
        try:
            if compressed:
                start = fc // 2
                raw = word_stream[start : start + char_count]
                piece = raw.decode("cp1252")
            else:
                raw = word_stream[fc : fc + char_count * 2]
                piece = raw.decode("utf-16le")
        except UnicodeDecodeError as exc:
            raise ExtractionError(
                "Failed to decode text from the uploaded .doc piece table."
            ) from exc
        pieces.append(piece)

    text = "".join(pieces)
    text = (
        text.replace("\x07", " ")
        .replace("\x0b", "\n")
        .replace("\x0c", "\n")
        .replace("\r", "\n")
    )
    return text


def _find_plcfpcd(clx: bytes) -> bytes | None:
    i = 0
    while i < len(clx):
        block_type = clx[i]
        i += 1
        if block_type == 0x01:
            if i + 2 > len(clx):
                return None
            cb = int.from_bytes(clx[i : i + 2], "little")
            i += 2 + cb
        elif block_type == 0x02:
            if i + 4 > len(clx):
                return None
            lcb = int.from_bytes(clx[i : i + 4], "little")
            i += 4
            if i + lcb > len(clx):
                return None
            return clx[i : i + lcb]
        else:
            return None
    return None


def _log_extraction_summary(document: ExtractedDocument) -> None:
    char_counts = [len(page.text) for page in document.pages]
    logger.info(
        "Extracted document metadata document_id=%s filename=%s file_type=%s "
        "page_count=%s total_chars=%s chars_per_page=%s",
        document.document_id,
        document.source_filename,
        document.file_type,
        document.page_count,
        sum(char_counts),
        char_counts,
    )


def load_url(
    url: str,
    document_id: str,
    timeout: float = 10.0,
) -> ExtractedDocument:
    """Fetch and extract readable text from a web URL."""
    try:
        import httpx
        from bs4 import BeautifulSoup
    except ImportError as e:
        raise ExtractionError(f"Missing dependency for URL extraction: {e}") from e

    try:
        response = httpx.get(
            url,
            timeout=timeout,
            follow_redirects=True,
            headers={"User-Agent": "TrustAwareRAG/1.0 (Research Ingestion Bot)"},
        )
        if response.status_code != 200:
            raise ExtractionError(f"Failed to fetch URL '{url}': HTTP {response.status_code}")
    except httpx.RequestError as exc:
        raise ExtractionError(f"Network error while fetching URL '{url}': {exc}") from exc

    html_content = response.text
    if not html_content.strip():
        raise ExtractionError(f"No content returned from URL: {url}")

    soup = BeautifulSoup(html_content, "html.parser")
    for element in soup(["script", "style", "nav", "footer", "header", "noscript", "svg"]):
        element.decompose()

    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    clean_text = "\n\n".join(lines)
    if not clean_text:
        raise ExtractionError(f"Could not extract readable text from URL: {url}")

    page = ExtractedPage(
        document_id=document_id,
        source_filename=url,
        page_number=1,
        text=clean_text,
    )
    doc = ExtractedDocument(
        document_id=document_id,
        source_filename=url,
        file_type="url",
        pages=(page,),
    )
    _log_extraction_summary(doc)
    return doc


def load_url_file(path: Path, document_id: str, filename: str) -> list[ExtractedPage]:
    """Extract URL from a .url shortcut file and fetch content."""
    raw = path.read_text(encoding="utf-8", errors="ignore")
    target_url = None
    for line in raw.splitlines():
        line = line.strip()
        if line.lower().startswith("url="):
            target_url = line.split("=", 1)[1].strip()
            break
        elif line.startswith("http://") or line.startswith("https://"):
            target_url = line
            break

    if not target_url:
        raise ExtractionError(f"No valid URL found in '{filename}'")

    doc = load_url(target_url, document_id=document_id)
    return list(doc.pages)


_LOADERS: dict[str, Callable[[Path, str, str], list[ExtractedPage]]] = {
    ".pdf": load_pdf,
    ".txt": load_txt,
    ".doc": load_doc,
    ".docx": load_docx,
    ".csv": load_csv,
    ".json": load_json,
    ".url": load_url_file,
}
